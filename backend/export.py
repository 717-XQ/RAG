# -*- coding: utf-8 -*-
"""
export.py - 文档导出模块（技术栈2.1：WeasyPrint HTML转PDF / python-docx生成.docx）
=============================================
功能：
  1. 将问答记录（问题/回答/引用来源）渲染为报告
  2. PDF导出：WeasyPrint 将HTML转PDF（Windows需安装GTK3运行时）
  3. Word导出：python-docx 生成 .docx 文档
  4. 中文兼容：CSS指定系统中文字体

使用方式：
  from export import export_pdf, export_docx
  pdf_bytes = export_pdf(question, answer, sources)
  docx_bytes = export_docx(question, answer, sources)
"""

import io
from typing import List, Dict, Any
from datetime import datetime

from loguru import logger

from config import settings


def _build_report_html(
    question: str,
    answer: str,
    sources: List[Dict[str, Any]],
    title: str = "",
) -> str:
    """
    构建报告HTML（WeasyPrint渲染用）

    Args:
        question: 用户问题
        answer: 回答文本
        sources: 引用来源列表 [{filename, page, score, content_preview}]
        title: 报告标题

    Returns:
        str: 完整HTML
    """
    title = title or settings.export.report_title
    font = settings.export.pdf_css_font

    source_items = ""
    if sources:
        items = []
        for i, s in enumerate(sources, 1):
            preview = (s.get("content_preview") or "")[:120]
            score = s.get("score", 0)
            items.append(
                f'<li><b>[{i}] {s.get("filename", "")}</b>'
                f'（第{s.get("page", 0)}页，相关度 {score:.2f}）'
                f'<br/><span class="preview">{preview}</span></li>'
            )
        source_items = '<div class="section-title">引用来源</div><ol class="sources">' + "".join(items) + "</ol>"
    else:
        source_items = '<p class="none">无引用来源（未检索到相关资料）</p>'

    # 回答文本转HTML（简单转义 + 换行）
    import html as html_lib
    answer_html = "<br/>".join(html_lib.escape(line) for line in answer.split("\n"))

    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<style>
  body {{ font-family: "{font}", "Microsoft YaHei", sans-serif; font-size: 12pt; color: #1e293b; line-height: 1.6; margin: 40px; }}
  h1 {{ font-size: 20pt; color: #1e3a8a; border-bottom: 2px solid #3b82f6; padding-bottom: 8px; }}
  .meta {{ color: #64748b; font-size: 9pt; margin-bottom: 20px; }}
  .section-title {{ font-size: 13pt; font-weight: bold; color: #1e3a8a; margin-top: 18px; border-left: 4px solid #3b82f6; padding-left: 8px; }}
  .question {{ background: #eff6ff; padding: 10px 14px; border-radius: 6px; }}
  .answer {{ white-space: normal; }}
  ol.sources li {{ margin-bottom: 8px; }}
  .preview {{ color: #64748b; font-size: 10pt; }}
  .none {{ color: #94a3b8; }}
  .footer {{ margin-top: 30px; color: #94a3b8; font-size: 8.5pt; text-align: center; }}
</style>
</head>
<body>
<h1>{html_lib.escape(title)}</h1>
<div class="meta">生成时间：{now}</div>

<div class="section-title">问题</div>
<div class="question">{html_lib.escape(question)}</div>

<div class="section-title">回答</div>
<div class="answer">{answer_html}</div>

{source_items}

<div class="footer">由 RAG 智能文档问答系统生成</div>
</body>
</html>"""


def export_pdf(
    question: str,
    answer: str,
    sources: List[Dict[str, Any]],
    title: str = "",
) -> bytes:
    """
    WeasyPrint 导出PDF（HTML转PDF）

    Args:
        question: 问题
        answer: 回答
        sources: 引用来源
        title: 报告标题

    Returns:
        bytes: PDF文件字节

    Raises:
        RuntimeError: Windows缺少GTK3运行时导致无法渲染
    """
    try:
        from weasyprint import HTML
    except ImportError:
        raise RuntimeError("未安装 weasyprint，请执行: pip install weasyprint")

    html_content = _build_report_html(question, answer, sources, title)

    try:
        pdf_bytes = HTML(string=html_content).write_pdf()
        logger.info(f"PDF导出成功: {len(pdf_bytes)} bytes")
        return pdf_bytes
    except Exception as e:
        err = str(e)
        if "libgobject" in err or "cannot load library" in err or "OSError" in err:
            raise RuntimeError(
                "PDF导出需要 GTK3 运行时。请安装 GTK3 Runtime for Windows "
                "（如 https://github.com/nicedash/gtk3-runtime/releases 或 "
                "https://github.com/tschoonj/GTK-for-Windows-Runtime-Environment-Installer），"
                "安装后重启终端即可。"
            )
        raise RuntimeError(f"PDF渲染失败: {err}")


def export_docx(
    question: str,
    answer: str,
    sources: List[Dict[str, Any]],
    title: str = "",
) -> bytes:
    """
    python-docx 导出Word文档

    Args:
        question: 问题
        answer: 回答
        sources: 引用来源
        title: 报告标题

    Returns:
        bytes: docx文件字节
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError:
        raise RuntimeError("未安装 python-docx，请执行: pip install python-docx")

    title = title or settings.export.report_title

    doc = Document()

    # 标题
    doc.add_heading(title, level=0)
    meta = doc.add_paragraph()
    run = meta.add_run(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # 问题
    doc.add_heading("问题", level=1)
    doc.add_paragraph(question)

    # 回答
    doc.add_heading("回答", level=1)
    doc.add_paragraph(answer)

    # 来源
    doc.add_heading("引用来源", level=1)
    if sources:
        for i, s in enumerate(sources, 1):
            p = doc.add_paragraph()
            p.add_run(f"[{i}] {s.get('filename', '')}（第{s.get('page', 0)}页，相关度 {float(s.get('score', 0)):.2f}）")
            preview = doc.add_paragraph((s.get("content_preview") or "")[:120])
            preview.runs[0].font.size = Pt(9)
            preview.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    else:
        doc.add_paragraph("无引用来源（未检索到相关资料）")

    # 页脚
    footer = doc.add_paragraph()
    run = footer.add_run("由 RAG 智能文档问答系统生成")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    data = buf.getvalue()
    logger.info(f"DOCX导出成功: {len(data)} bytes")
    return data
